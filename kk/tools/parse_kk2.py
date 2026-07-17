# -*- coding: utf-8 -*-
import json, re, csv, zipfile, io, os
import openpyxl, xlrd

PREFS=[("北海道","01"),("青森","02"),("岩手","03"),("宮城","04"),("秋田","05"),("山形","06"),("福島","07"),("茨城","08"),("栃木","09"),("群馬","10"),("埼玉","11"),("千葉","12"),("東京","13"),("神奈川","14"),("新潟","15"),("富山","16"),("石川","17"),("福井","18"),("山梨","19"),("長野","20"),("岐阜","21"),("静岡","22"),("愛知","23"),("三重","24"),("滋賀","25"),("京都","26"),("大阪","27"),("兵庫","28"),("奈良","29"),("和歌山","30"),("鳥取","31"),("島根","32"),("岡山","33"),("広島","34"),("山口","35"),("徳島","36"),("香川","37"),("愛媛","38"),("高知","39"),("福岡","40"),("佐賀","41"),("長崎","42"),("熊本","43"),("大分","44"),("宮崎","45"),("鹿児島","46"),("沖縄","47")]
ROMAJI={"hokkaido":"01","aomori":"02","iwate":"03","miyagi":"04","akita":"05","yamagata":"06","fukushima":"07","ibaraki":"08","tochigi":"09","gunma":"10","saitama":"11","chiba":"12","tokyo":"13","kanagawa":"14","niigata":"15","toyama":"16","ishikawa":"17","fukui":"18","yamanashi":"19","nagano":"20","gifu":"21","shizuoka":"22","aichi":"23","mie":"24","shiga":"25","kyoto":"26","osaka":"27","hyogo":"28","nara":"29","wakayama":"30","tottori":"31","shimane":"32","okayama":"33","hiroshima":"34","yamaguchi":"35","tokushima":"36","kagawa":"37","ehime":"38","kochi":"39","fukuoka":"40","saga":"41","nagasaki":"42","kumamoto":"43","oita":"44","miyazaki":"45","kagoshima":"46","okinawa":"47"}
_RS=sorted(ROMAJI.items(), key=lambda x:-len(x[0]))
def pref_of(t):
    for nm,cd in PREFS:
        if nm in t: return cd
    tl=t.lower()
    for nm,cd in _RS:
        if nm in tl: return cd
    return None
def mt_of(t):
    cands=[("医科",t.find("医科")),("歯科",t.find("歯科")),("薬局",t.find("薬局"))]
    hit=[c for c in cands if c[1]>=0]
    if hit: return sorted(hit,key=lambda x:x[1])[0][0]
    tl=t.lower()
    si=tl.find("shika"); ii=tl.find("ika"); yi=tl.find("yakkyoku")
    if si>=0 and (ii<0 or si<=ii<=si+2): ii=-1
    c2=[("医科",ii),("歯科",si),("薬局",yi)]
    h2=[c for c in c2 if c[1]>=0]
    return sorted(h2,key=lambda x:x[1])[0][0] if h2 else None
DG=re.compile(r"[^0-9]")
PHONE=re.compile(r"^[0-9０-９]{2,5}[-−ー－][0-9０-９]")
BED=re.compile(r"(一般|療養|精神|結核|感染|全床)\D{0,6}?([0-9０-９,，]+)")
ZIPC=re.compile(r"^〒?[0-9０-９]{3}[－ー\-−][0-9０-９]{4}")

def sheets_of(name, data):
    if name.lower().endswith(".xls"):
        bk=xlrd.open_workbook(file_contents=data)
        for sn in bk.sheet_names():
            sh=bk.sheet_by_name(sn)
            yield sn, [[("" if sh.cell_value(r,c) is None else str(sh.cell_value(r,c)).strip()) for c in range(sh.ncols)] for r in range(sh.nrows)]
    else:
        wb=openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sn in wb.sheetnames:
            ws=wb[sn]
            yield sn, [[("" if v is None else str(v).strip()) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()

def parse_sheet(rows, pref, mt, src):
    hr=-1
    for i in range(min(len(rows),30)):
        j="".join(rows[i])
        if ("医療機関" in j or "薬局" in j) and ("番号" in j or "コード" in j) and len(j)<200:
            hr=i; break
    ci=None; start=None
    if hr>=0:
        nxt=rows[hr+1] if hr+1<len(rows) else []
        comb=[(rows[hr][i] if i<len(rows[hr]) else "")+" "+(nxt[i] if i<len(nxt) else "") for i in range(max(len(rows[hr]),len(nxt)))]
        def find(pat):
            for i,h in enumerate(comb):
                if re.search(pat,h): return i
            return -1
        ci=dict(code=find(r"(医療機関|薬局).*(番号|コード)"), name=find(r"名\s*称"), addr=find(r"所在地"),
                phone=find(r"電話"), opener=find(r"開設者"), manager=find(r"管理者"),
                depts=find(r"診療科"), desig=find(r"指定年月日"), kind=-1)
        if ci["code"]<0 or ci["name"]<0: ci=None
        else: start=hr+2
    if ci is None:
        for i in range(min(len(rows),45)):
            r=rows[i]
            if len(r)>2 and len(DG.sub("",r[1] if len(r)>1 else ""))>=7 and (r[2].strip() if len(r)>2 else ""):
                ci=dict(code=1,name=2,addr=3,phone=4,opener=5,manager=6,desig=7,depts=8,kind=9)
                start=i; break
        if ci is None: return None, "no_cols"
    recs={}; cur=[None]
    def g(row,k):
        i=ci.get(k,-1)
        return row[i].strip() if 0<=i<len(row) else ""
    def flush():
        c=cur[0]
        if c and c["med_code"]:
            b=0
            for lab,nu in BED.findall(c.pop("_bedsrc","")):
                try: b+=int(DG.sub("",nu) or 0)
                except: pass
            c["beds"]=str(b) if b>0 else ""
            dep=c.get("depts","")
            c["depts"]="　".join(t for t in re.split(r"[\s　、，,]+",dep) if t and not re.search(r"[0-9０-９]",t))[:500]
            recs[c["med_code"]]=c
        cur[0]=None
    for r in range(start,len(rows)):
        row=rows[r]
        dg=DG.sub("",g(row,"code"))
        nm=g(row,"name")
        if len(dg)>=7 and nm:
            flush()
            c=dict(med_type=mt,pref_code=pref,med_code=dg[-7:],name=nm[:120],
                   address=ZIPC.sub("",g(row,"addr"))[:300],
                   phone=g(row,"phone")[:40] if PHONE.match(g(row,"phone")) else "",
                   opener=g(row,"opener")[:200],manager=g(row,"manager")[:120],
                   depts=g(row,"depts"),_bedsrc=g(row,"depts"),
                   designated_date=g(row,"desig")[:20],source=src)
            k=g(row,"kind")
            if mt=="医科" and k in ("病院","診療所"): c["med_type"]=k
            cur[0]=c
        elif cur[0]:
            c=cur[0]
            a=g(row,"addr"); d=g(row,"depts"); p=g(row,"phone"); n2=g(row,"name"); op=g(row,"opener"); mg=g(row,"manager")
            if a: c["address"]=(c["address"]+a)[:300]
            if d:
                c["depts"]=c["depts"]+"　"+d
                c["_bedsrc"]=c.get("_bedsrc","")+"　"+d
            if p and not c["phone"] and PHONE.match(p): c["phone"]=p[:40]
            if n2 and not dg and len(c["name"])<60: c["name"]=(c["name"]+n2)[:120]
            if op and len(c["opener"])<180: c["opener"]=(c["opener"]+op)[:200]
            if mg and not c["manager"]: c["manager"]=mg[:120]
    flush()
    return list(recs.values()), None

def main():
    man=json.load(open("/home/claude/kk/manifest.json"))
    out={}; stats={}; skips=[]; fails=[]
    for m in man:
        fp="/home/claude/kk/files/"+os.path.basename(m["path"])
        if not os.path.exists(fp): fails.append((m["id"],"missing")); continue
        data=open(fp,"rb").read()
        entries=[]
        if data[:2]==b"PK" and fp.lower().endswith(".zip"):
            try:
                z=zipfile.ZipFile(io.BytesIO(data))
                for zi in z.namelist():
                    if re.search(r"\.(xlsx|xls)$", zi, re.I): entries.append((zi, z.read(zi)))
            except Exception as e:
                fails.append((m["id"],"zip:"+str(e)[:60])); continue
        else:
            entries.append((os.path.basename(fp), data))
        for en,ed in entries:
            try:
                for sn,rows in sheets_of(en,ed):
                    ctx=en+" "+sn+" "+(m.get("link_text") or "")
                    pref=pref_of(ctx) or m.get("pref")
                    mt=mt_of(ctx) or m.get("mt")
                    if not pref or not mt: skips.append((m["id"],en[:30],sn[:15],"class")); continue
                    if len(rows)<3: continue
                    recs,err=parse_sheet(rows,pref,mt,m["url"][:120])
                    if err: skips.append((m["id"],en[:30],sn[:15],err)); continue
                    out.setdefault(m["bureau"],{})
                    for rec in recs:
                        out[m["bureau"]][(rec["med_type"],rec["pref_code"],rec["med_code"])]=rec
                    stats[(m["bureau"],pref,mt)]=stats.get((m["bureau"],pref,mt),0)+len(recs)
            except Exception as e:
                fails.append((m["id"],en[:40]+":"+str(e)[:60]))
    cols=["med_type","pref_code","med_code","name","address","phone","opener","manager","depts","beds","designated_date","source"]
    for b,d in out.items():
        with open("/home/claude/kk/out/%s.csv"%b,"w",newline="",encoding="utf-8") as f:
            w=csv.DictWriter(f,fieldnames=cols)
            w.writeheader()
            for rec in d.values():
                rec.pop("_bedsrc",None)
                w.writerow(rec)
    json.dump({"stats":{"%s|%s|%s"%k:v for k,v in sorted(stats.items())},
               "totals":{b:len(d) for b,d in out.items()},
               "skips":skips[:40],"n_skips":len(skips),"fails":fails[:40],"n_fails":len(fails)},
              open("/home/claude/kk/out/_report.json","w"),ensure_ascii=False,indent=1)
    print("bureaus:",{b:len(d) for b,d in out.items()},"skips:",len(skips),"fails:",len(fails))

if __name__=="__main__": main()
